from __future__ import annotations

import argparse
import hashlib
import re
import sqlite3
import time
from dataclasses import dataclass
from datetime import datetime, timedelta
from pathlib import Path
from typing import Callable

from openpyxl import load_workbook


WEEKDAYS = [
    "monday",
    "tuesday",
    "wednesday",
    "thursday",
    "friday",
    "saturday",
    "sunday",
]

EYE_ALIASES = {
    "left": "left",
    "l": "left",
    "le": "left",
    "right": "right",
    "r": "right",
    "re": "right",
    "both": "both",
    "ou": "both",
    "bilateral": "both",
}

TIME_12H_RE = re.compile(r"^\s*(\d{1,2}):(\d{2})\s*([AaPp][Mm])\s*$")
TIME_24H_RE = re.compile(r"^\s*([01]?\d|2[0-3]):([0-5]\d)\s*$")


@dataclass(frozen=True)
class Reminder:
    medicine: str
    eye: str
    weekday: str
    time_hhmm: str


def canonical_text(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def normalize_time(time_text: str) -> str:
    text = canonical_text(time_text)
    m12 = TIME_12H_RE.match(text)
    if m12:
        hour = int(m12.group(1)) % 12
        minute = int(m12.group(2))
        mer = m12.group(3).lower()
        if mer == "pm":
            hour += 12
        return f"{hour:02d}:{minute:02d}"

    m24 = TIME_24H_RE.match(text)
    if m24:
        return f"{int(m24.group(1)):02d}:{int(m24.group(2)):02d}"

    raise ValueError(f"Unrecognized time format: {time_text!r}")


def format_time_display(hhmm: str) -> str:
    h, m = map(int, hhmm.split(":"))
    dt = datetime(2000, 1, 1, h, m)
    return dt.strftime("%-I:%M %p")


def normalize_eye(eye_text: str) -> str:
    key = canonical_text(eye_text).lower()
    normalized = EYE_ALIASES.get(key)
    if not normalized:
        raise ValueError(f"Unrecognized eye column: {eye_text!r}")
    return normalized


def stable_weekday(text: str) -> str:
    key = canonical_text(text).lower()
    if key not in WEEKDAYS:
        raise ValueError(f"Invalid weekday: {text!r}")
    return key


def detect_headers(sheet) -> tuple[int, int, dict[int, str], dict[int, str]]:
    """
    Detect: header row index, time column index, weekday columns, eye columns.

    Expected layout (flexible):
    - One row with weekday names spread across columns.
    - Optional next row with eye labels under each weekday column.
    - One column with time values in data rows.
    """

    max_row = sheet.max_row
    max_col = sheet.max_column

    header_row = None
    weekday_cols: dict[int, str] = {}

    for r in range(1, max_row + 1):
        found = {}
        for c in range(1, max_col + 1):
            val = canonical_text(sheet.cell(r, c).value).lower()
            if val in WEEKDAYS:
                found[c] = val
        if len(found) >= 2:
            header_row = r
            weekday_cols = found
            break

    if header_row is None:
        raise RuntimeError("Could not detect weekday header row.")

    eye_cols: dict[int, str] = {}
    eye_row = header_row + 1
    for c in weekday_cols.keys():
        v = canonical_text(sheet.cell(eye_row, c).value)
        if v:
            eye_cols[c] = normalize_eye(v)

    if not eye_cols:
        for c in weekday_cols.keys():
            eye_cols[c] = "both"

    time_col = None
    for c in range(1, max_col + 1):
        values = [canonical_text(sheet.cell(r, c).value) for r in range(header_row + 1, min(max_row + 1, header_row + 20))]
        ok = 0
        for v in values:
            if not v:
                continue
            try:
                normalize_time(v)
                ok += 1
            except ValueError:
                pass
        if ok >= 2:
            time_col = c
            break

    if time_col is None:
        raise RuntimeError("Could not detect time column.")

    return header_row, time_col, weekday_cols, eye_cols


def parse_sheet_to_reminders(workbook_path: Path, sheet_name: str | None = None) -> list[Reminder]:
    wb = load_workbook(workbook_path, data_only=True)
    sheet = wb[sheet_name] if sheet_name else wb.active

    header_row, time_col, weekday_cols, eye_cols = detect_headers(sheet)

    reminders: list[Reminder] = []
    for r in range(header_row + 1, sheet.max_row + 1):
        raw_time = canonical_text(sheet.cell(r, time_col).value)
        if not raw_time:
            continue
        try:
            hhmm = normalize_time(raw_time)
        except ValueError:
            continue

        for c, weekday in weekday_cols.items():
            if c == time_col:
                continue
            medicine = canonical_text(sheet.cell(r, c).value)
            if not medicine:
                continue
            reminders.append(
                Reminder(
                    medicine=medicine,
                    eye=eye_cols.get(c, "both"),
                    weekday=stable_weekday(weekday),
                    time_hhmm=hhmm,
                )
            )

    unique = {(x.medicine, x.eye, x.weekday, x.time_hhmm): x for x in reminders}
    return sorted(unique.values(), key=lambda x: (WEEKDAYS.index(x.weekday), x.time_hhmm, x.eye, x.medicine.lower()))


def file_hash(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        while True:
            chunk = f.read(1024 * 64)
            if not chunk:
                break
            h.update(chunk)
    return h.hexdigest()


def init_db(conn: sqlite3.Connection) -> None:
    conn.executescript(
        """
        CREATE TABLE IF NOT EXISTS worksheet_state (
            id INTEGER PRIMARY KEY CHECK (id = 1),
            workbook_hash TEXT NOT NULL,
            updated_at TEXT NOT NULL
        );

        CREATE TABLE IF NOT EXISTS reminders (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            medicine TEXT NOT NULL,
            eye TEXT NOT NULL CHECK (eye IN ('left', 'right', 'both')),
            weekday TEXT NOT NULL,
            time_hhmm TEXT NOT NULL,
            active INTEGER NOT NULL DEFAULT 1,
            UNIQUE (medicine, eye, weekday, time_hhmm)
        );

        CREATE TABLE IF NOT EXISTS reminder_log (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            reminder_id INTEGER NOT NULL,
            due_date TEXT NOT NULL,
            sent_at TEXT NOT NULL,
            UNIQUE (reminder_id, due_date),
            FOREIGN KEY (reminder_id) REFERENCES reminders(id)
        );
        """
    )
    conn.commit()


def refresh_db_from_sheet(conn: sqlite3.Connection, workbook_path: Path, sheet_name: str | None = None) -> bool:
    current_hash = file_hash(workbook_path)
    row = conn.execute("SELECT workbook_hash FROM worksheet_state WHERE id = 1").fetchone()
    if row and row[0] == current_hash:
        return False

    reminders = parse_sheet_to_reminders(workbook_path, sheet_name)

    conn.execute("DELETE FROM reminders")
    conn.executemany(
        "INSERT INTO reminders (medicine, eye, weekday, time_hhmm, active) VALUES (?, ?, ?, ?, 1)",
        [(r.medicine, r.eye, r.weekday, r.time_hhmm) for r in reminders],
    )

    conn.execute(
        """
        INSERT INTO worksheet_state (id, workbook_hash, updated_at)
        VALUES (1, ?, ?)
        ON CONFLICT(id) DO UPDATE SET workbook_hash=excluded.workbook_hash, updated_at=excluded.updated_at
        """,
        (current_hash, datetime.utcnow().isoformat()),
    )
    conn.commit()
    return True


def send_due_reminders(conn: sqlite3.Connection, now: datetime, sender: Callable[[str], None]) -> int:
    target = now + timedelta(minutes=5)
    target_weekday = WEEKDAYS[target.weekday()]
    target_hhmm = target.strftime("%H:%M")
    target_date = target.date().isoformat()

    rows = conn.execute(
        """
        SELECT id, medicine, eye, time_hhmm
        FROM reminders
        WHERE active = 1 AND weekday = ? AND time_hhmm = ?
        """,
        (target_weekday, target_hhmm),
    ).fetchall()

    sent = 0
    for reminder_id, medicine, eye, hhmm in rows:
        exists = conn.execute(
            "SELECT 1 FROM reminder_log WHERE reminder_id = ? AND due_date = ?",
            (reminder_id, target_date),
        ).fetchone()
        if exists:
            continue

        eye_phrase = "both eyes" if eye == "both" else f"your {eye} eye"
        msg = f"Reminder: In 5 minutes, take {medicine} in {eye_phrase} at {format_time_display(hhmm)}."
        sender(msg)

        conn.execute(
            "INSERT INTO reminder_log (reminder_id, due_date, sent_at) VALUES (?, ?, ?)",
            (reminder_id, target_date, now.isoformat()),
        )
        sent += 1

    conn.commit()
    return sent


def print_sender(message: str) -> None:
    print(message)


def run_loop(db_path: Path, workbook_path: Path, sheet_name: str | None, interval_seconds: int) -> None:
    conn = sqlite3.connect(db_path)
    init_db(conn)

    while True:
        changed = refresh_db_from_sheet(conn, workbook_path, sheet_name)
        if changed:
            print("Worksheet update detected. Reminders refreshed.")

        sent = send_due_reminders(conn, now=datetime.utcnow(), sender=print_sender)
        if sent:
            print(f"Sent {sent} reminder(s).")

        time.sleep(interval_seconds)


def main() -> None:
    parser = argparse.ArgumentParser(description="Eye medicine weekly reminder engine.")
    parser.add_argument("workbook", type=Path, help="Path to .xlsx worksheet")
    parser.add_argument("--sheet", type=str, default=None, help="Worksheet name (default: active sheet)")
    parser.add_argument("--db", type=Path, default=Path("reminders.db"), help="SQLite DB path")
    parser.add_argument("--run-once", action="store_true", help="Only refresh DB from sheet and exit")
    parser.add_argument("--tick", type=int, default=60, help="Loop interval seconds (default: 60)")
    args = parser.parse_args()

    conn = sqlite3.connect(args.db)
    init_db(conn)

    changed = refresh_db_from_sheet(conn, args.workbook, args.sheet)
    if changed:
        print("Worksheet parsed and reminders stored.")
    else:
        print("No worksheet changes detected.")

    if args.run_once:
        return

    run_loop(args.db, args.workbook, args.sheet, args.tick)


if __name__ == "__main__":
    main()
